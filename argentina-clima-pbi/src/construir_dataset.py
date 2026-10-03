"""
construir_dataset.py
======================

Construcción de los datasets nacional y provincial a partir de los
archivos crudos de data/raw/. No se imputan valores faltantes ni se
inventan datos: donde una fuente no cubre un año/jurisdicción, queda NaN.

Ejecutar:
    python -m src.construir_dataset
"""

import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

from src.data_loader import PROVINCIAS_ARGENTINA, CABA, NACIONAL

RAW_DIR = Path(__file__).resolve().parent.parent / "data" / "raw"
PROCESSED_DIR = Path(__file__).resolve().parent.parent / "data" / "processed"
TABLES_DIR = Path(__file__).resolve().parent.parent / "outputs" / "tables"

PERIODO_BASE_ANOMALIA = (1991, 2020)


def _normalizar(texto: str) -> str:
    texto = texto.replace("_", " ").strip()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn"
    )
    return texto.lower().strip()


_JURISDICCIONES_CANONICAS = PROVINCIAS_ARGENTINA + [CABA]
_INDICE_NORMALIZADO = {_normalizar(j): j for j in _JURISDICCIONES_CANONICAS}
_ALIAS = {
    "ciudad autonoma de buenos aires": CABA,
    "ciudad de buenos aires": CABA,
    "caba": CABA,
    "pba": "Buenos Aires",
    "tierra del fuego": "Tierra del Fuego, Antártida e Islas del Atlántico Sur",
    "tierra del fuego antartida e islas del atlantico sur": (
        "Tierra del Fuego, Antártida e Islas del Atlántico Sur"
    ),
}
for alias, canonico in _ALIAS.items():
    _INDICE_NORMALIZADO[_normalizar(alias)] = canonico


def mapear_jurisdiccion(nombre_crudo: str) -> str | None:
    if nombre_crudo is None:
        return None
    return _INDICE_NORMALIZADO.get(_normalizar(str(nombre_crudo)))


ESTACION_A_PROVINCIA = {
    "Buenos Aires": CABA,
    "La Plata": "Buenos Aires",
    "Catamarca": "Catamarca",
    "Córdoba": "Córdoba",
    "Corrientes": "Corrientes",
    "Formosa": "Formosa",
    "Jujuy": "Jujuy",
    "La Rioja": "La Rioja",
    "Mendoza": "Mendoza",
    "Neuquén": "Neuquén",
    "Paraná": "Entre Ríos",
    "Posadas": "Misiones",
    "Resistencia": "Chaco",
    "Río Gallegos": "Santa Cruz",
    "Salta": "Salta",
    "San Juan": "San Juan",
    "San Luis": "San Luis",
    "Santa Rosa": "La Pampa",
    "Santiago del Estero": "Santiago del Estero",
    "Sauce Viejo": "Santa Fe",
    "Trelew": "Chubut",
    "Tucumán": "Tucumán",
    "Ushuaia": "Tierra del Fuego, Antártida e Islas del Atlántico Sur",
    "Viedma": "Río Negro",
}
ESTACIONES_NO_UTILIZADAS = {"Mar del Plata", "Rosario"}


def cargar_anomalia_nacional() -> pd.DataFrame:
    df = pd.read_csv(RAW_DIR / "ciam_anomalia_temp_precip_nacional.csv", sep=";")
    df.columns = ["año", "anomalia_temperatura", "anomalia_precipitacion_pct"]
    df["jurisdiccion"] = NACIONAL
    df["tipo_jurisdiccion"] = "nacional"
    return df[["año", "jurisdiccion", "tipo_jurisdiccion", "anomalia_temperatura",
               "anomalia_precipitacion_pct"]]


def cargar_precipitacion_provincial() -> pd.DataFrame:
    df = pd.read_csv(RAW_DIR / "ciam_precipitacion_estaciones.csv", sep=";")
    df = df.rename(columns={df.columns[0]: "estacion"})
    df["estacion"] = df["estacion"].str.strip()
    df = df[~df["estacion"].isin(ESTACIONES_NO_UTILIZADAS)]
    df["jurisdiccion"] = df["estacion"].map(ESTACION_A_PROVINCIA)
    sin_mapear = df[df["jurisdiccion"].isna()]
    if not sin_mapear.empty:
        print(f"AVISO: estaciones sin mapear a provincia: {sin_mapear['estacion'].tolist()}")

    columnas_año = [c for c in df.columns if c.startswith("año_")]
    largo = df.melt(
        id_vars=["jurisdiccion"], value_vars=columnas_año,
        var_name="año", value_name="precipitaciones_mm",
    )
    largo["año"] = largo["año"].str.replace("año_", "").astype(int)
    largo["precipitaciones_mm"] = (
        largo["precipitaciones_mm"].astype(str).str.strip().astype(float)
    )
    largo["tipo_jurisdiccion"] = largo["jurisdiccion"].apply(
        lambda j: "CABA" if j == CABA else "provincia"
    )
    return largo.dropna(subset=["jurisdiccion"])


def cargar_temperatura_provincial() -> pd.DataFrame:
    df = pd.read_csv(RAW_DIR / "nasa_power_temperatura_provincial.csv")
    df.columns = ["jurisdiccion", "año", "temperatura_media_c", "latitud", "longitud"]
    df["jurisdiccion"] = df["jurisdiccion"].apply(mapear_jurisdiccion)
    df.loc[df["temperatura_media_c"] <= -900, "temperatura_media_c"] = pd.NA

    inicio, fin = PERIODO_BASE_ANOMALIA
    base = df[(df["año"] >= inicio) & (df["año"] <= fin)]
    medias_base = base.groupby("jurisdiccion")["temperatura_media_c"].mean()
    df["temperatura_media_base_1991_2020"] = df["jurisdiccion"].map(medias_base)
    df["anomalia_temperatura"] = df["temperatura_media_c"] - df["temperatura_media_base_1991_2020"]

    df["tipo_jurisdiccion"] = df["jurisdiccion"].apply(
        lambda j: "CABA" if j == CABA else "provincia"
    )
    return df[["año", "jurisdiccion", "tipo_jurisdiccion", "temperatura_media_c",
               "anomalia_temperatura"]]


def cargar_gei_provincial_y_nacional() -> pd.DataFrame:
    wb = openpyxl.load_workbook(
        RAW_DIR / "desagregacion-provincial_hasta_2022.xlsx", read_only=True
    )
    filas = []
    for nombre_hoja in wb.sheetnames:
        if nombre_hoja == "0_Sin Asignar":
            continue
        ws = wb[nombre_hoja]
        filas_hoja = list(ws.iter_rows(values_only=True))

        fila_años = None
        for fila in filas_hoja:
            if fila[0] is not None and isinstance(fila[1], (int, float)) and fila[1] and 1990 < fila[1] < 2100:
                fila_años = fila
                break
        if fila_años is None:
            print(f"AVISO: no se encontró fila de años en la hoja '{nombre_hoja}'")
            continue
        años = [int(a) for a in fila_años[1:] if a is not None]

        fila_total = None
        for fila in filas_hoja:
            etiqueta = str(fila[0]).strip().lower() if fila[0] else ""
            if etiqueta.startswith("total pais") or etiqueta.startswith("total jurisdiccion"):
                fila_total = fila
                break
        if fila_total is None:
            print(f"AVISO: no se encontró fila de total en la hoja '{nombre_hoja}'")
            continue

        if nombre_hoja == "Total Pais":
            jurisdiccion = NACIONAL
        else:
            nombre_sin_codigo = nombre_hoja.split("_", 1)[-1] if "_" in nombre_hoja else nombre_hoja
            jurisdiccion = mapear_jurisdiccion(nombre_sin_codigo)
            if jurisdiccion is None:
                print(f"AVISO: hoja '{nombre_hoja}' no matchea con ninguna jurisdicción")
                continue

        valores = fila_total[1:1 + len(años)]
        for año, valor in zip(años, valores):
            filas.append({"año": año, "jurisdiccion": jurisdiccion, "emisiones_gei_co2eq": valor})

    df = pd.DataFrame(filas)
    df["tipo_jurisdiccion"] = df["jurisdiccion"].apply(
        lambda j: "nacional" if j == NACIONAL else ("CABA" if j == CABA else "provincia")
    )
    return df


def cargar_co2_fosil_nacional() -> pd.DataFrame:
    df = pd.read_csv(RAW_DIR / "owid_co2_data_argentina.csv")
    df = df[["year", "co2"]].rename(columns={"year": "año", "co2": "emisiones_co2_fosil"})
    df["jurisdiccion"] = NACIONAL
    df["tipo_jurisdiccion"] = "nacional"
    return df


def cargar_pbi_nacional() -> pd.DataFrame:
    df = pd.read_csv(RAW_DIR / "banco_mundial_pbi_crecimiento.csv")
    df["jurisdiccion"] = NACIONAL
    df["tipo_jurisdiccion"] = "nacional"
    return df


def cargar_poblacion_pbi_percapita_nacional() -> pd.DataFrame:
    """Población y PBI (Maddison Project Database, vía OWID) para calcular
    PBI per cápita nacional. Mismo archivo crudo que las emisiones de CO2
    fósil (owid_co2_data_argentina.csv), pero es una variable distinta:
    documentada aparte en data/sources.csv."""
    df = pd.read_csv(RAW_DIR / "owid_co2_data_argentina.csv")
    df = df[["year", "population", "gdp"]].rename(
        columns={"year": "año", "population": "poblacion"}
    )
    df["pib_per_capita"] = df["gdp"] / df["poblacion"]
    df = df.sort_values("año").reset_index(drop=True)
    año_anterior = df["año"].shift(1)
    valor_anterior = df["pib_per_capita"].shift(1)
    consecutivo = (df["año"] - año_anterior) == 1
    df["pib_per_capita_crecimiento_pct"] = (
        (df["pib_per_capita"] / valor_anterior - 1) * 100
    ).where(consecutivo)
    return df[["año", "poblacion", "pib_per_capita", "pib_per_capita_crecimiento_pct"]]


def _parsear_año_encabezado(valor) -> int | None:
    if isinstance(valor, (int, float)):
        return int(valor)
    if isinstance(valor, str):
        digitos = "".join(c for c in valor[:4] if c.isdigit())
        if len(digitos) == 4:
            return int(digitos)
    return None


def cargar_pbg_provincial() -> pd.DataFrame:
    wb = openpyxl.load_workbook(RAW_DIR / "cepal_pbg_provincial.xlsx", read_only=True)
    ws = wb["VABpb"]
    filas_hoja = list(ws.iter_rows(values_only=True))

    fila_encabezado = None
    for fila in filas_hoja:
        if fila[1] == "JURISDICCIÓN":
            fila_encabezado = fila
            break
    if fila_encabezado is None:
        raise ValueError("No se encontró la fila de encabezado 'JURISDICCIÓN' en la hoja VABpb.")

    años_crudos = fila_encabezado[2:]
    años = [_parsear_año_encabezado(a) for a in años_crudos]
    es_preliminar = [isinstance(a, str) and "(2)" in a for a in años_crudos]
    años = [a for a in años if a is not None]

    idx_inicio = filas_hoja.index(fila_encabezado) + 1
    filas = []
    for fila in filas_hoja[idx_inicio:]:
        nombre_crudo = fila[1]
        if not nombre_crudo or str(nombre_crudo).strip() in ("Total", "No distribuido"):
            continue
        jurisdiccion = mapear_jurisdiccion(nombre_crudo)
        if jurisdiccion is None:
            continue
        valores = fila[2:2 + len(años)]
        for año, valor, preliminar in zip(años, valores, es_preliminar):
            filas.append({
                "año": año, "jurisdiccion": jurisdiccion, "pbg": valor,
                "pbg_preliminar": preliminar,
            })

    df = pd.DataFrame(filas)
    df["tipo_jurisdiccion"] = df["jurisdiccion"].apply(
        lambda j: "CABA" if j == CABA else "provincia"
    )
    return df


def construir_dataset_nacional() -> pd.DataFrame:
    anomalia = cargar_anomalia_nacional()
    co2eq = cargar_gei_provincial_y_nacional()
    co2eq_nacional = co2eq[co2eq["jurisdiccion"] == NACIONAL][["año", "emisiones_gei_co2eq"]]
    co2_fosil = cargar_co2_fosil_nacional()[["año", "emisiones_co2_fosil"]]
    pbi = cargar_pbi_nacional()[["año", "pbi_crecimiento_pct"]]
    percapita = cargar_poblacion_pbi_percapita_nacional()

    n_filas_base = len(anomalia)
    df = anomalia.merge(co2eq_nacional, on="año", how="left")
    df = df.merge(co2_fosil, on="año", how="left")
    df = df.merge(pbi, on="año", how="left")
    df = df.merge(percapita, on="año", how="left")
    if len(df) > n_filas_base:
        raise ValueError("El merge del dataset nacional multiplicó filas.")
    return df.sort_values("año").reset_index(drop=True)


def construir_dataset_provincial() -> pd.DataFrame:
    temperatura = cargar_temperatura_provincial()
    precipitacion = cargar_precipitacion_provincial()
    emisiones = cargar_gei_provincial_y_nacional()
    emisiones_prov = emisiones[emisiones["jurisdiccion"] != NACIONAL][
        ["año", "jurisdiccion", "tipo_jurisdiccion", "emisiones_gei_co2eq"]
    ]
    pbg = cargar_pbg_provincial()[["año", "jurisdiccion", "pbg", "pbg_preliminar"]]

    df = temperatura.merge(
        precipitacion[["año", "jurisdiccion", "precipitaciones_mm"]],
        on=["año", "jurisdiccion"], how="outer",
    )
    df = df.merge(
        emisiones_prov[["año", "jurisdiccion", "emisiones_gei_co2eq"]],
        on=["año", "jurisdiccion"], how="outer",
    )
    df = df.merge(pbg, on=["año", "jurisdiccion"], how="outer")

    df["tipo_jurisdiccion"] = df["jurisdiccion"].apply(
        lambda j: "CABA" if j == CABA else "provincia"
    )

    duplicados = df.duplicated(subset=["año", "jurisdiccion"]).sum()
    if duplicados > 0:
        raise ValueError(f"El dataset provincial quedó con {duplicados} filas duplicadas.")

    columnas = ["año", "jurisdiccion", "tipo_jurisdiccion", "temperatura_media_c",
                "anomalia_temperatura", "precipitaciones_mm", "emisiones_gei_co2eq",
                "pbg", "pbg_preliminar"]
    return df[columnas].sort_values(["jurisdiccion", "año"]).reset_index(drop=True)


def generar_informe_cobertura(df: pd.DataFrame, nombre: str) -> pd.DataFrame:
    filas = []
    columnas_variables = [
        c for c in df.columns if c not in ("año", "jurisdiccion", "tipo_jurisdiccion")
    ]
    for col in columnas_variables:
        no_nulos = df[df[col].notna()]
        filas.append({
            "dataset": nombre,
            "variable": col,
            "primer_año": no_nulos["año"].min() if not no_nulos.empty else None,
            "ultimo_año": no_nulos["año"].max() if not no_nulos.empty else None,
            "observaciones": no_nulos.shape[0],
            "faltantes": df[col].isna().sum(),
            "jurisdicciones_con_datos": (
                no_nulos["jurisdiccion"].nunique() if "jurisdiccion" in df.columns else None
            ),
        })
    return pd.DataFrame(filas)


def main():
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    TABLES_DIR.mkdir(parents=True, exist_ok=True)

    print("Construyendo dataset nacional...")
    nacional = construir_dataset_nacional()
    nacional.to_csv(PROCESSED_DIR / "dataset_nacional.csv", index=False)
    print(f"  -> {len(nacional)} filas, {nacional['año'].min()}-{nacional['año'].max()}")

    print("Construyendo dataset provincial...")
    provincial = construir_dataset_provincial()
    provincial.to_csv(PROCESSED_DIR / "dataset_provincial.csv", index=False)
    print(f"  -> {len(provincial)} filas, {provincial['jurisdiccion'].nunique()} jurisdicciones")

    print("Generando informe de cobertura...")
    informe = pd.concat([
        generar_informe_cobertura(nacional, "nacional"),
        generar_informe_cobertura(provincial, "provincial"),
    ])
    informe.to_csv(TABLES_DIR / "informe_cobertura.csv", index=False)
    print(informe.to_string(index=False))


if __name__ == "__main__":
    main()
